"""Generate KB Image Extraction Approach Word document for Lily/team review."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    def add_heading_styled(text, level=1):
        p = doc.add_paragraph()
        run = p.add_run(text)
        if level == 1:
            set_run_font(run, 16, True, (0, 51, 102))
            p.space_before = Pt(14)
            p.space_after = Pt(8)
        elif level == 2:
            set_run_font(run, 13, True, (0, 70, 120))
            p.space_before = Pt(12)
            p.space_after = Pt(6)
        else:
            set_run_font(run, 11, True, (40, 40, 40))
            p.space_before = Pt(8)
            p.space_after = Pt(4)
        return p

    def add_body(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_run_font(run, 11)
        p.paragraph_format.space_after = Pt(6)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style="List Bullet")
        if bold_prefix:
            r1 = p.add_run(bold_prefix)
            set_run_font(r1, 11, True)
            r2 = p.add_run(text)
            set_run_font(r2, 11)
        else:
            r = p.add_run(text)
            set_run_font(r, 11)
        p.paragraph_format.space_after = Pt(3)
        return p

    def add_numbered(text):
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(text)
        set_run_font(r, 11)
        p.paragraph_format.space_after = Pt(3)
        return p

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Image Extraction Approach for SpyderWash Knowledge Base")
    set_run_font(r, 18, True, (0, 51, 102))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run(
        "How we will extract, clean, store, and retrieve images from the manuals"
    )
    set_run_font(r, 11, False, (80, 80, 80))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Prepared for: Lily Susanna / Team Review\nDate: 1 August 2026")
    set_run_font(r, 10, False, (100, 100, 100))

    # 1
    add_heading_styled("1. Why do we need image extraction?", 1)
    add_body(
        "Our Operator AI answers technical questions using the Knowledge Base text. "
        "But many troubleshooting steps are easier to follow when the operator can also see "
        "the related screenshot, diagram, or device photo."
    )
    add_body(
        "Right now, those images sit inside Word documents. The AI cannot use them directly. "
        "So we need a clear process to pull images out of the manuals, connect them to the "
        "correct support cases, and let the agent retrieve them when needed."
    )

    # 2
    add_heading_styled("2. Which documents will we use?", 1)
    add_body("We will work with the two active Knowledge Base files:")
    add_bullet(
        "structured support cases (main source first)",
        "SpyderWash AI Support Knowledge Base (v2.2).docx — ",
    )
    add_bullet(
        "troubleshooting guide and operator FAQ (after the v2.2 pipeline is stable)",
        "Setomatic Bible.docx — ",
    )
    add_body(
        "Both files already contain embedded images (screenshots, diagrams, photos). "
        "v2.2 has about 66 images. The Bible has about 220 images."
    )

    # 3
    add_heading_styled(
        "3. Chosen approach: Convert DOCX to PDF first (Flow A)", 1
    )
    add_body(
        "We will first convert each Word document into a PDF, then extract images from the PDF. "
        "This is better than pulling images straight from the Word file."
    )
    add_heading_styled("Why PDF is better for this work", 2)
    add_bullet(
        "Page numbers become stable and easy to quote (for example: “see page 12”)."
    )
    add_bullet(
        "Tables and diagrams stay together as one visual page, instead of being split "
        "across Word XML pieces."
    )
    add_bullet(
        "Some Bible pages are scanned photos. Full-page PDF rendering captures those correctly."
    )
    add_bullet(
        "Mapping an image to a support case becomes simpler using page ranges."
    )
    add_bullet(
        "PyMuPDF can extract embedded images and also render full pages when needed."
    )

    add_heading_styled("What we will not do", 2)
    add_body(
        "We will not extract images only by unzipping the Word file. "
        "That method loses reliable page order and makes case-to-image mapping harder to audit."
    )

    # 4
    add_heading_styled("4. High-level flow (simple view)", 1)
    add_body("End-to-end flow in plain words:")
    add_numbered("Convert the Word manual to PDF.")
    add_numbered(
        "Remove the first six pages (as requested), so processed page 1 = original page 7."
    )
    add_numbered("Skip empty pages. We will not store blank or near-blank pages.")
    add_numbered(
        "Find which pages belong to which support case (using ARTICLE START / ARTICLE END boundaries)."
    )
    add_numbered("Extract screenshots and diagrams from each useful page.")
    add_numbered(
        "If an embedded image is incomplete (for example a cut table), render the full page at high resolution."
    )
    add_numbered(
        "Remove duplicate logos, headers, icons, and tiny decorative images."
    )
    add_numbered(
        "Generate a short caption and visual summary for each useful image."
    )
    add_numbered(
        "Store searchable terms (button names, error messages, device labels, etc.)."
    )
    add_numbered(
        "Link images to the matching case IDs, and also to linked companion cases."
    )
    add_numbered(
        "Expose retrieval tools so the Operator Agent can fetch the right images with the answer."
    )
    add_numbered("Log what was retrieved and run test cases.")

    # 5
    add_heading_styled("5. Step-by-step explanation", 1)

    add_heading_styled("Step 1 — Prepare the source PDF", 2)
    add_body(
        "We convert the DOCX file into PDF. Then we remove the first six pages. "
        "This keeps a clean page offset so processed page 1 always matches original page 7. "
        "This makes page references consistent for the whole team."
    )

    add_heading_styled("Step 2 — Skip empty pages", 2)
    add_body(
        "Lily’s guidance is clear: do not store empty pages. "
        "A page is treated as empty when it has little or no text, no useful images, "
        "and almost no visual content. We skip those pages completely — no image extraction, "
        "no storage, no database entry."
    )

    add_heading_styled("Step 3 — Map pages to support cases", 2)
    add_body(
        "Each v2.2 article already has clear ARTICLE START and ARTICLE END markers. "
        "We use those markers to build a page range for every case ID "
        "(for example: KB-READER-001 covers pages 12–15). "
        "Any useful image found in that range is linked to that case."
    )

    add_heading_styled("Step 4 — Extract images", 2)
    add_body("For every useful page, we do two things:")
    add_bullet(
        "First try to pull out the embedded screenshot or diagram from the PDF.",
        "Embedded extraction: ",
    )
    add_bullet(
        "If the embedded image does not keep the full screenshot, table, or diagram, "
        "we render the whole page at high resolution and save that as the visual.",
        "Full-page render (fallback): ",
    )

    add_heading_styled("Step 5 — Remove duplicates and junk", 2)
    add_body("We clean the image set before storage:")
    add_bullet("Same image appearing many times is kept only once (checksum match).")
    add_bullet("Very small images (logos, icons) are dropped.")
    add_bullet("Header/footer style strips and decorative graphics are dropped.")
    add_bullet("Invalid or broken image files are dropped.")

    add_heading_styled("Step 6 — Caption and search terms", 2)
    add_body(
        "For each clean image, we generate a short caption and a visual summary. "
        "The summary captures what an operator would notice: buttons, fields, labels, "
        "error messages, arrows, and device state. "
        "From that text, we also store search terms so images can be found later by UI label, "
        "device name, product, or action."
    )

    add_heading_styled(
        "Step 7 — Link images to cases (including linked cases)", 2
    )
    add_body(
        "Images are linked to the active case first. "
        "If that case references another case (for example CASE-007 → CASE-010), "
        "we also return images from the linked case, in the same order as the procedure. "
        "This reuses our existing co-retrieval / case-reference graph."
    )

    add_heading_styled("Step 8 — Retrieval tools and agent use", 2)
    add_body("We will expose simple tools such as:")
    add_bullet("get_case_images — return images for one case ID")
    add_bullet("search_case_images — find images by search term / device / type")
    add_bullet(
        "get_case_visual_bundle — return text case + related images together"
    )
    add_body(
        "The Operator Agent prompt will be updated so it retrieves case text and images together, "
        "follows linked-case order, and does not invent visual details that are not in the image."
    )

    # 6
    add_heading_styled("6. Why these tools?", 1)
    add_bullet(
        "Reliable PDF reading, page rendering, and image extraction.",
        "PyMuPDF: ",
    )
    add_bullet(
        "Same structured store we already use for articles and co-retrieval rules.",
        "SQLite: ",
    )
    add_bullet(
        "Checksum-based duplicate detection and clean metadata "
        "(path, page, caption, type).",
        "Image database tables: ",
    )
    add_bullet(
        "Accurate captions and UI search terms from screenshots.",
        "Vision model (for captions): ",
    )
    add_bullet(
        "Same style as our existing KB search tools.",
        "MCP / API tools: ",
    )

    # 7
    add_heading_styled("7. What will be stored for each image?", 1)
    add_body("For every useful image we keep:")
    add_bullet("File path and standardized filename")
    add_bullet("Linked case ID(s)")
    add_bullet("Page number (processed page and original page)")
    add_bullet("Image type (screenshot, diagram, table, photo, warning)")
    add_bullet("Caption and visual summary")
    add_bullet("Dimensions, checksum, and sequence order")
    add_bullet("Search terms for later retrieval")

    add_heading_styled("What we will not store", 2)
    add_bullet("Empty pages")
    add_bullet("Duplicate logos / headers / decorative icons")
    add_bullet("Tiny or invalid images")
    add_bullet("Internal-only decorative graphics with no operator value")

    # 8
    add_heading_styled("8. Benefits of this approach", 1)
    add_bullet(
        "Operators get the right screenshot with the right troubleshooting steps."
    )
    add_bullet(
        "Images stay tied to case IDs, so retrieval is controlled and auditable."
    )
    add_bullet(
        "Linked cases can return images in procedure order, not randomly."
    )
    add_bullet(
        "Empty and junk visuals are filtered out early, so the database stays clean."
    )
    add_bullet(
        "The same pipeline can later be reused for the Bible document."
    )

    # 9
    add_heading_styled("9. Delivery plan", 1)
    add_body(
        "To keep quality high, we will first complete the full pipeline on the v2.2 Knowledge Base "
        "(cleaner structure, fewer images). After that works end to end, we will run the same "
        "pipeline on the Setomatic Bible."
    )
    add_body("Planned order:")
    add_numbered(
        "Requirements, PDF prep, schema, extraction, dedup, and storage"
    )
    add_numbered(
        "Captions, search terms, case linking, and retrieval queries"
    )
    add_numbered(
        "MCP tools, agent prompt update, logging, test cases, and end-to-end checks"
    )

    # 10
    add_heading_styled("10. Expected outcome", 1)
    add_body(
        "At the end of this work, the Operator AI will be able to answer with grounded case text "
        "and the matching visuals — screenshots, diagrams, and UI instructions — without showing "
        "internal article IDs or inventing image details that are not present."
    )

    close = doc.add_paragraph()
    r = close.add_run(
        "Happy to walk through this approach on a short call if needed. "
        "Please share any feedback so we can lock the flow before full implementation."
    )
    set_run_font(r, 11)

    sign = doc.add_paragraph()
    r = sign.add_run("\nRegards,\nMonish Pilla")
    set_run_font(r, 11)

    out_path = os.path.join("docs", "KB_Image_Extraction_Approach.docx")
    os.makedirs("docs", exist_ok=True)
    doc.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
