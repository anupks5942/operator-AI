"""Parse v2.3 Image Asset Registry table and store in SQLite.

Brandon's v2.3 KB includes a table (typically the last table in the DOCX)
with columns: Image ID | Description | Related article(s).

This module:
1. Parses that table from the DOCX.
2. Stores rows in `image_asset_registry` SQLite table.
3. Provides lookup helpers so the image pipeline can use Brandon's captions
   and article mappings instead of GPT-4o generated ones.
"""
from __future__ import annotations

import logging
import os
import sqlite3
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

_DB_PATH = os.path.normpath(
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "kb_structured.db")
)

_REGISTRY_TABLE_HEADER = "image id"


@dataclass
class AssetRow:
    image_id: str
    description: str
    related_articles: list[str] = field(default_factory=list)


def parse_asset_registry(docx_path: str) -> list[AssetRow]:
    """Parse the Image Asset Registry table from a v2.3+ DOCX file."""
    from docx import Document

    doc = Document(docx_path)
    registry_table = None

    for table in doc.tables:
        if not table.rows:
            continue
        header_text = table.rows[0].cells[0].text.strip().lower()
        if _REGISTRY_TABLE_HEADER in header_text:
            registry_table = table
            break

    if registry_table is None:
        for table in reversed(doc.tables):
            if not table.rows:
                continue
            if len(table.columns) == 3:
                first_data = table.rows[1].cells[0].text.strip() if len(table.rows) > 1 else ""
                if first_data.startswith("IMG-SW-"):
                    registry_table = table
                    break

    if registry_table is None:
        logger.warning("[ASSET REGISTRY] No Image Asset Registry table found in %s", docx_path)
        return []

    rows: list[AssetRow] = []
    for row in registry_table.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if len(cells) < 3:
            continue
        image_id = cells[0].strip()
        description = cells[1].strip()
        related_raw = cells[2].strip()

        if not image_id.startswith("IMG-SW-"):
            continue

        related_articles = [
            a.strip() for a in related_raw.replace("\n", ",").split(",") if a.strip()
        ]

        rows.append(AssetRow(
            image_id=image_id,
            description=description,
            related_articles=related_articles,
        ))

    logger.info("[ASSET REGISTRY] Parsed %d image entries from %s", len(rows), docx_path)
    return rows


def create_registry_table(db_path: str | None = None) -> None:
    """Create the image_asset_registry table if it doesn't exist."""
    db = db_path or _DB_PATH
    conn = sqlite3.connect(db)
    try:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS image_asset_registry (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                image_id TEXT UNIQUE NOT NULL,
                description TEXT NOT NULL DEFAULT '',
                source_doc TEXT NOT NULL DEFAULT 'v2.3',
                version TEXT NOT NULL DEFAULT '2.3',
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS image_registry_articles (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                image_id TEXT NOT NULL,
                article_id TEXT NOT NULL,
                UNIQUE(image_id, article_id)
            );

            CREATE INDEX IF NOT EXISTS idx_registry_image_id
                ON image_asset_registry(image_id);
            CREATE INDEX IF NOT EXISTS idx_registry_articles_image
                ON image_registry_articles(image_id);
            CREATE INDEX IF NOT EXISTS idx_registry_articles_article
                ON image_registry_articles(article_id);
        """)
        conn.commit()
    finally:
        conn.close()


def ingest_registry(rows: list[AssetRow], db_path: str | None = None, source_doc: str = "v2.3") -> int:
    """Insert parsed registry rows into SQLite. Returns count inserted."""
    db = db_path or _DB_PATH
    create_registry_table(db)

    conn = sqlite3.connect(db)
    count = 0
    try:
        for row in rows:
            conn.execute(
                """INSERT OR REPLACE INTO image_asset_registry
                   (image_id, description, source_doc, version)
                   VALUES (?, ?, ?, ?)""",
                (row.image_id, row.description, source_doc, "2.3"),
            )
            for article_id in row.related_articles:
                conn.execute(
                    """INSERT OR IGNORE INTO image_registry_articles
                       (image_id, article_id) VALUES (?, ?)""",
                    (row.image_id, article_id),
                )
            count += 1
        conn.commit()
        logger.info("[ASSET REGISTRY] Ingested %d registry entries into %s", count, db)
    finally:
        conn.close()
    return count


def get_registry_articles_for_image(image_id: str, db_path: str | None = None) -> list[str]:
    """Return article IDs linked to an image from the registry."""
    db = db_path or _DB_PATH
    conn = sqlite3.connect(db)
    try:
        rows = conn.execute(
            "SELECT article_id FROM image_registry_articles WHERE image_id = ?",
            (image_id,),
        ).fetchall()
        return [r[0] for r in rows]
    finally:
        conn.close()


def get_registry_images_for_article(article_id: str, db_path: str | None = None) -> list[dict]:
    """Return all registry images linked to an article."""
    db = db_path or _DB_PATH
    conn = sqlite3.connect(db)
    try:
        rows = conn.execute("""
            SELECT r.image_id, r.description
            FROM image_asset_registry r
            JOIN image_registry_articles ra ON r.image_id = ra.image_id
            WHERE ra.article_id = ?
            ORDER BY r.image_id
        """, (article_id,)).fetchall()
        return [{"image_id": r[0], "description": r[1]} for r in rows]
    finally:
        conn.close()


def get_registry_description(image_id: str, db_path: str | None = None) -> str:
    """Return Brandon's description for an image ID."""
    db = db_path or _DB_PATH
    conn = sqlite3.connect(db)
    try:
        row = conn.execute(
            "SELECT description FROM image_asset_registry WHERE image_id = ?",
            (image_id,),
        ).fetchone()
        return row[0] if row else ""
    finally:
        conn.close()
